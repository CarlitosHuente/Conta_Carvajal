"""Libro auxiliar de remuneraciones (art. 62 del Código del Trabajo)."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .liquidacion_items import descuentos_para_presentacion, items_monto_suma
from .models import Liquidacion


MESES = (
    '', 'Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
    'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre',
)

# Mismos bloques que pide la Dirección del Trabajo para el libro:
# identificación, haberes, descuentos, aportes del empleador y totales.
COLUMNAS = (
    ('rut', 'RUT', 'Identificación', False),
    ('nombre', 'Trabajador', 'Identificación', False),
    ('cargo', 'Cargo', 'Identificación', False),
    ('ingreso', 'Fecha ingreso', 'Identificación', False),
    ('dias', 'Días trab.', 'Identificación', False),
    ('afp', 'AFP', 'Identificación', False),
    ('salud', 'Salud', 'Identificación', False),
    ('sueldo_base', 'Sueldo base', 'Haberes', True),
    ('gratificacion', 'Gratificación', 'Haberes', True),
    ('horas_extra', 'Horas extra', 'Haberes', True),
    ('otros_imponibles', 'Otros imponibles', 'Haberes', True),
    ('total_imponible', 'Total imponible', 'Haberes', True),
    ('colacion', 'Colación', 'Haberes', True),
    ('movilizacion', 'Movilización', 'Haberes', True),
    ('asignacion_familiar', 'Asig. familiar', 'Haberes', True),
    ('otros_no_imponibles', 'Otros no imponibles', 'Haberes', True),
    ('total_no_imponible', 'Total no imponible', 'Haberes', True),
    ('total_haberes', 'Total haberes', 'Haberes', True),
    ('dcto_afp', 'AFP trabajador', 'Descuentos', True),
    ('dcto_salud', 'Salud 7%', 'Descuentos', True),
    ('dcto_adicional', 'Adicional Isapre', 'Descuentos', True),
    ('dcto_cesantia', 'Cesantía trabajador', 'Descuentos', True),
    ('dcto_impuesto', 'Impuesto único', 'Descuentos', True),
    ('dcto_otros', 'Otros descuentos', 'Descuentos', True),
    ('total_descuentos', 'Total descuentos', 'Descuentos', True),
    ('aporte_sis', 'SIS empleador', 'Aportes empleador', True),
    ('aporte_afc', 'AFC empleador', 'Aportes empleador', True),
    ('total_aportes', 'Total aportes empleador', 'Aportes empleador', True),
    ('liquido', 'Líquido a pagar', 'Totales', True),
)

_CLAVES_MONTO = [clave for clave, _etiq, _grupo, es_monto in COLUMNAS if es_monto]


def _suma_prefijo(items, prefijos):
    return items_monto_suma(items, prefijos)


def fila_liquidacion(liq):
    haberes = [i for i in liq.items.all() if i.tipo == 'HABER']
    descuentos = descuentos_para_presentacion(liq)
    sueldo = _suma_prefijo(haberes, ['Sueldo Base'])
    grat = _suma_prefijo(haberes, ['Gratificación', 'Gratificacion'])
    horas = _suma_prefijo(haberes, ['Horas Extra', 'Horas extra'])
    colacion = _suma_prefijo(haberes, ['Colación', 'Colacion'])
    movilizacion = _suma_prefijo(haberes, ['Movilización', 'Movilizacion'])
    asig = liq.total_asignacion_familiar or _suma_prefijo(haberes, ['Asignación Familiar', 'Asignacion Familiar'])
    imponible = liq.total_haberes_imponibles
    no_imponible = liq.total_haberes_no_imponibles
    otros_imp = max(0, imponible - sueldo - grat - horas)
    otros_no = max(0, no_imponible - colacion - movilizacion - asig)

    dcto_afp = _suma_prefijo(descuentos, ['AFP '])
    dcto_salud = _suma_prefijo(descuentos, ['Salud '])
    dcto_adicional = _suma_prefijo(descuentos, ['Adicional Isapre'])
    dcto_cesantia = _suma_prefijo(descuentos, ['Seguro de Cesant'])
    dcto_impuesto = _suma_prefijo(descuentos, ['Impuesto '])
    total_descuentos = liq.total_descuentos_legales + liq.total_descuentos_varios
    dcto_otros = max(0, total_descuentos - dcto_afp - dcto_salud - dcto_adicional - dcto_cesantia - dcto_impuesto)
    total_aportes = liq.cotizacion_sis_empleador + liq.cotizacion_afc_empleador
    trabajador = liq.contrato.trabajador
    ingreso = liq.fecha_ingreso_contrato or liq.contrato.fecha_inicio

    return {
        'rut': trabajador.rut,
        'nombre': trabajador.nombre_completo,
        'cargo': liq.cargo_contrato or liq.contrato.cargo or '',
        'ingreso': ingreso.strftime('%d/%m/%Y') if ingreso else '',
        'dias': liq.dias_trabajados,
        'afp': liq.afp_nombre or liq.contrato.afp.nombre,
        'salud': liq.salud_nombre or liq.contrato.sistema_salud.nombre,
        'sueldo_base': sueldo,
        'gratificacion': grat,
        'horas_extra': horas,
        'otros_imponibles': otros_imp,
        'total_imponible': imponible,
        'colacion': colacion,
        'movilizacion': movilizacion,
        'asignacion_familiar': asig,
        'otros_no_imponibles': otros_no,
        'total_no_imponible': no_imponible,
        'total_haberes': imponible + no_imponible,
        'dcto_afp': dcto_afp,
        'dcto_salud': dcto_salud,
        'dcto_adicional': dcto_adicional,
        'dcto_cesantia': dcto_cesantia,
        'dcto_impuesto': dcto_impuesto,
        'dcto_otros': dcto_otros,
        'total_descuentos': total_descuentos,
        'aporte_sis': liq.cotizacion_sis_empleador,
        'aporte_afc': liq.cotizacion_afc_empleador,
        'total_aportes': total_aportes,
        'liquido': liq.sueldo_liquido,
    }


def libro_del_periodo(empresa, mes, ano):
    liquidaciones = (
        Liquidacion.objects.filter(contrato__trabajador__empresa=empresa, mes=mes, ano=ano)
        .select_related('contrato__trabajador', 'contrato__afp', 'contrato__sistema_salud')
        .prefetch_related('items')
        .order_by('contrato__trabajador__apellido_paterno', 'contrato__trabajador__apellido_materno')
    )
    filas = [fila_liquidacion(liq) for liq in liquidaciones]
    totales = {clave: sum(fila[clave] for fila in filas) for clave in _CLAVES_MONTO}
    return filas, totales


def excel_libro(empresa, mes, ano, filas, totales):
    wb = Workbook()
    ws = wb.active
    ws.title = f'{mes:02d}-{ano}'

    titulo = Font(bold=True, size=14)
    ws['A1'] = 'LIBRO AUXILIAR DE REMUNERACIONES'
    ws['A1'].font = titulo
    ws['A2'] = empresa.razon_social
    ws['A3'] = f'RUT empleador: {empresa.rut}'
    ws['A4'] = f'Período: {MESES[mes]} {ano}'
    ws['A5'] = (
        'Registro del artículo 62 del Código del Trabajo. '
        'La declaración ante la Dirección del Trabajo se hace en el Libro de Remuneraciones Electrónico (archivo CSV del portal DT).'
    )
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=8)

    header_fill = {
        'Identificación': PatternFill('solid', fgColor='1F4E79'),
        'Haberes': PatternFill('solid', fgColor='1E7A46'),
        'Descuentos': PatternFill('solid', fgColor='8C2F2F'),
        'Aportes empleador': PatternFill('solid', fgColor='7A5B1E'),
        'Totales': PatternFill('solid', fgColor='0F6B4C'),
    }
    blanco = Font(bold=True, color='FFFFFF', size=9)
    borde = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0'),
    )
    fila_grupo = 7
    fila_titulo = 8
    col = 1
    grupo_actual = None
    inicio_grupo = 1
    for clave, etiqueta, grupo, _es_monto in COLUMNAS:
        if grupo != grupo_actual:
            if grupo_actual is not None and col - 1 > inicio_grupo:
                ws.merge_cells(start_row=fila_grupo, start_column=inicio_grupo, end_row=fila_grupo, end_column=col - 1)
            grupo_actual = grupo
            inicio_grupo = col
            celda = ws.cell(fila_grupo, col, grupo)
            celda.fill = header_fill[grupo]
            celda.font = blanco
            celda.alignment = Alignment(horizontal='center')
        else:
            celda = ws.cell(fila_grupo, col, '')
            celda.fill = header_fill[grupo]
        titulo_celda = ws.cell(fila_titulo, col, etiqueta)
        titulo_celda.fill = header_fill[grupo]
        titulo_celda.font = blanco
        titulo_celda.alignment = Alignment(horizontal='center', wrap_text=True)
        titulo_celda.border = borde
        col += 1
    if col - 1 > inicio_grupo:
        ws.merge_cells(start_row=fila_grupo, start_column=inicio_grupo, end_row=fila_grupo, end_column=col - 1)

    fila_excel = 9
    for fila in filas:
        for indice, (clave, _etiq, _grupo, es_monto) in enumerate(COLUMNAS, start=1):
            celda = ws.cell(fila_excel, indice, fila[clave])
            celda.border = borde
            if es_monto:
                celda.number_format = '#,##0'
                celda.alignment = Alignment(horizontal='right')
        fila_excel += 1

    if filas:
        ws.cell(fila_excel, 1, 'TOTALES').font = Font(bold=True)
        for indice, (clave, _etiq, _grupo, es_monto) in enumerate(COLUMNAS, start=1):
            celda = ws.cell(fila_excel, indice)
            celda.font = Font(bold=True)
            celda.border = borde
            celda.fill = PatternFill('solid', fgColor='E7E6E6')
            if es_monto:
                celda.value = totales[clave]
                celda.number_format = '#,##0'
        ws.merge_cells(start_row=fila_excel, start_column=1, end_row=fila_excel, end_column=7)

    for indice, (_clave, _etiq, _grupo, es_monto) in enumerate(COLUMNAS, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = 16 if es_monto else 18
    ws.column_dimensions['B'].width = 32
    ws.row_dimensions[fila_titulo].height = 30
    ws.freeze_panes = 'A9'
    ws.auto_filter.ref = f'A8:{get_column_letter(len(COLUMNAS))}{max(fila_excel - 1, 8)}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = empresa.razon_social
    ws.oddFooter.left.text = 'Libro auxiliar de remuneraciones — art. 62 Código del Trabajo'
    ws.oddFooter.right.text = f'{MESES[mes]} {ano}'

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
